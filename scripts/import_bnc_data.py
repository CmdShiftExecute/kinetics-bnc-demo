"""Import the three supplied BNC workbooks into one deterministic private JSON file.

Only project-level and company-level fields are retained. Contact names, telephone
numbers, email addresses and workbook assignee fields are deliberately not imported.
Money remains in the source currency and is stored as USD million to one decimal.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SOURCES = (
    {
        "key": "urban_industrial",
        "file": "3-BNC-Urban-Ind.-Greenfield-Projects-Master-File.xlsx",
        "sheet": "Combined Projects",
        "header": 1,
        "reference": "BNC Reference",
        "snapshot": "2026-01-01",
    },
    {
        "key": "other_sectors",
        "file": "2-BNC_Other_Sectors_Master_File.xlsx",
        "sheet": "Combined Projects",
        "header": 1,
        "reference": "BNC Reference",
        "snapshot": "2026-03-01",
    },
    {
        "key": "brownfield",
        "file": "1-Raw-Data-BNC-Brownfield-All-Sectors.xlsx",
        "sheet": "Project",
        "header": 10,
        "reference": "Reference Number",
        "snapshot": "2026-05-04",
    },
)


def text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def iso(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    value = text(value)
    if not value or value.upper() in {"N/A", "NA", "NONE", "NULL"}:
        return None
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    raw = text(value).replace(",", "").replace("$", "")
    if not raw or raw.upper() in {"N/A", "NA", "NONE", "NULL", "-"}:
        return None
    try:
        return float(raw)
    except ValueError:
        return None


def stage(value: Any) -> str:
    raw = text(value)
    if raw.startswith("Concept"):
        return "Concept"
    if raw.startswith("Design"):
        return "Design"
    if raw.startswith("Tender"):
        return "Tender"
    if raw == "UC" or raw.startswith("Under Construction"):
        return "Under Construction"
    if raw.startswith("Completed"):
        return raw
    raise ValueError(f"Unrecognised BNC stage: {raw!r}")


def completion(value: Any, project_stage: str) -> float | None:
    if project_stage != "Under Construction":
        return None
    n = number(value)
    if n is None:
        return None
    if 0 <= n <= 1:
        n *= 100
    return round(max(0, min(100, n)), 1)


def attributes(value: Any) -> list[str]:
    result: list[str] = []
    for item in (text(part) for part in str(value or "").split(",")):
        if item and item not in result:
            result.append(item)
    return result


def one_company_field(value: Any) -> list[str]:
    """Preserve the workbook cell verbatim rather than guessing at ambiguous commas."""
    value = text(value)
    if value.casefold() in {"not yet awarded", "see sub-projects"}:
        return []
    return [value] if value else []


def rows(path: Path, source: dict[str, Any]) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[source["sheet"]]
    headers = [text(cell.value) for cell in sheet[source["header"]]]
    index = {header: i for i, header in enumerate(headers) if header}
    for values in sheet.iter_rows(min_row=source["header"] + 1, values_only=True):
        ref = text(values[index[source["reference"]]])
        if not ref:
            continue
        yield {header: values[i] for header, i in index.items()} | {"_ref": ref}


def project(row: dict[str, Any], source: dict[str, Any]) -> dict[str, Any]:
    project_stage = stage(row.get("Stage"))
    value_usd = number(row.get("Value (USD)", row.get("Value(USD)")))
    last_updated = iso(row.get("Last Updated", row.get("Updated Date"))) or source["snapshot"]
    category = text(row.get("Project Category")) or ("Brownfield" if source["key"] == "brownfield" else "Greenfield")
    return {
        "ref": row["_ref"],
        "name": text(row.get("Project Name")),
        "stage": project_stage,
        "completionPct": completion(row.get("Completion %", row.get("Completion Percentage")), project_stage),
        "completionDate": iso(row.get("Completion Date")),
        "valueUsd": value_usd,
        "value": round(value_usd / 1_000_000, 1) if value_usd is not None and value_usd > 0 else 0.0,
        "country": text(row.get("Country")),
        "city": text(row.get("City")),
        "sector": text(row.get("Sector")),
        "category": category,
        "industry": text(row.get("Industry")) or "Not recorded",
        "type": text(row.get("Project Type")) or "Not recorded",
        "location": text(row.get("Location")),
        "attributes": attributes(row.get("Project Attributes")),
        "owners": one_company_field(row.get("Owners")),
        "leadConsultants": one_company_field(row.get("Lead / Design Consultants")),
        "mepConsultants": one_company_field(row.get("MEP Consultants")),
        "mainContractors": one_company_field(row.get("Main / EPC Contractors")),
        "mepContractors": one_company_field(row.get("MEP Contractors")),
        "description": text(row.get("Description")),
        "lastUpdated": last_updated,
        "source": source["key"],
    }


def load_canonical(directory: Path) -> dict[str, Any]:
    combined: dict[str, dict[str, Any]] = {}
    source_meta = []
    raw_rows = 0
    for source in SOURCES:
        path = directory / source["file"]
        imported = [project(row, source) for row in rows(path, source)]
        raw_rows += len(imported)
        source_meta.append(
            {
                "key": source["key"],
                "file": source["file"],
                "sheet": source["sheet"],
                "snapshot": source["snapshot"],
                "rows": len(imported),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        )
        # Later, newer sources win on an overlapping BNC reference.
        for item in imported:
            combined[item["ref"]] = item
    projects = sorted(combined.values(), key=lambda item: item["ref"])
    return {
        "metadata": {
            "currency": "USD",
            "unit": "million",
            "rawRows": raw_rows,
            "uniqueProjects": len(projects),
            "overlapRows": raw_rows - len(projects),
            "sources": source_meta,
            "privacy": "Project and company fields only; contact names, phone numbers, emails and assignees excluded.",
        },
        "projects": projects,
    }


def main() -> None:
    root = Path(__file__).parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=root / ".private" / "bnc")
    parser.add_argument("--output", type=Path, default=root / ".private" / "bnc-source.json")
    args = parser.parse_args()
    data = load_canonical(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(
        f"Imported {data['metadata']['rawRows']:,} rows into "
        f"{data['metadata']['uniqueProjects']:,} unique BNC projects; "
        f"{data['metadata']['overlapRows']} overlaps resolved by source recency."
    )


if __name__ == "__main__":
    main()
